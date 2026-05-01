"""
Module 4b — Investment Memo (Python version)
=============================================
Generates a professional Word document investment memo.
Uses python-docx — no Node.js required.

Usage:
    pip install python-docx
    python 08_investment_memo.py
"""

import os
import pandas as pd
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Config ────────────────────────────────────────────────────────────────────
RISK_DIR  = "risk_output"
ML_DIR    = "ml_output"
OUT_DIR   = "dashboard_output"
os.makedirs(OUT_DIR, exist_ok=True)

BLUE      = RGBColor(0x15, 0x65, 0xC0)
DARK_BLUE = RGBColor(0x0D, 0x3C, 0x6E)
GRAY      = RGBColor(0x54, 0x6E, 0x7A)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x21, 0x21, 0x21)
TODAY     = datetime.today().strftime("%B %d, %Y")


# ── Load CSV data ─────────────────────────────────────────────────────────────

def load(path, index_col=None):
    if os.path.exists(path):
        return pd.read_csv(path, index_col=index_col)
    print(f"  WARNING: {path} not found — using placeholder data")
    return pd.DataFrame()


risk_df   = load(f"{RISK_DIR}/risk_metrics.csv",    index_col=0)
port_df   = load(f"{RISK_DIR}/portfolio_summary.csv")
var_df    = load(f"{RISK_DIR}/var_results.csv",     index_col=0)
ml_df     = load(f"{ML_DIR}/model_evaluation.csv")
feat_df   = load(f"{ML_DIR}/feature_importance.csv", index_col=0)

port      = port_df.iloc[0].to_dict() if not port_df.empty else {}
best_ml   = ml_df.sort_values("roc_auc", ascending=False).iloc[0].to_dict() \
            if not ml_df.empty else {}
top3      = risk_df.sort_values("sharpe", ascending=False).head(3) \
            if not risk_df.empty and "sharpe" in risk_df.columns else pd.DataFrame()
top_feats = feat_df.head(3) if not feat_df.empty else pd.DataFrame()


# ── Helper functions ──────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    """Set table cell background colour."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def set_cell_border(cell, hex_color="BBDEFB"):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"),   "single")
        border.set(qn("w:sz"),    "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), hex_color)
        tcBorders.append(border)
    tcPr.append(tcBorders)


def add_paragraph_border_bottom(para, hex_color="1565C0"):
    pPr  = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "6")
    bot.set(qn("w:space"), "4")
    bot.set(qn("w:color"), hex_color)
    pBdr.append(bot)
    pPr.append(pBdr)


def heading1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    add_paragraph_border_bottom(p)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(14)
    run.font.color.rgb = DARK_BLUE
    run.font.name = "Arial"
    return p


def heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    run.font.name = "Arial"
    return p


def body(doc, text, italic=False, color=None):
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text)
    run.font.size   = Pt(10)
    run.font.name   = "Arial"
    run.italic      = italic
    run.font.color.rgb = color if color else DARK
    return p


def bullet(doc, text):
    p   = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "Arial"
    run.font.color.rgb = DARK
    return p


def kpi_table(doc, rows_data):
    """Two-column label/value table for KPIs."""
    table = doc.add_table(rows=len(rows_data), cols=2)
    table.style = "Table Grid"
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.0)

    for i, (label, value) in enumerate(rows_data):
        row = table.rows[i]
        bg  = "E3F2FD" if i % 2 == 0 else "FFFFFF"

        c0, c1 = row.cells[0], row.cells[1]
        set_cell_bg(c0, bg); set_cell_border(c0)
        set_cell_bg(c1, bg); set_cell_border(c1)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.size = Pt(9); r0.font.name = "Arial"
        r0.font.color.rgb = GRAY

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1 = p1.add_run(str(value))
        r1.bold = True; r1.font.size = Pt(9)
        r1.font.name = "Arial"; r1.font.color.rgb = DARK_BLUE

    return table


def metrics_table(doc, headers, rows, col_widths_in=None):
    """Generic data table with blue header row."""
    n_cols = len(headers)
    table  = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.style = "Table Grid"

    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            for cell in table.columns[i].cells:
                cell.width = Inches(w)

    # Header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        set_cell_bg(c, "1565C0")
        set_cell_border(c, "0D3C6E")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True; r.font.size = Pt(9)
        r.font.name = "Arial"; r.font.color.rgb = WHITE

    # Data rows
    for ri, row_vals in enumerate(rows):
        row = table.rows[ri + 1]
        bg  = "F5F7FA" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row_vals):
            c = row.cells[ci]
            set_cell_bg(c, bg); set_cell_border(c)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.size = Pt(9); r.font.name = "Arial"
            r.font.color.rgb = DARK

    return table


def spacer(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)


# ── Build document ────────────────────────────────────────────────────────────

def build_memo():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # ── COVER PAGE ────────────────────────────────────────────────────────────
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run("EQUITY RESEARCH PLATFORM")
    r.bold = True; r.font.size = Pt(24)
    r.font.color.rgb = DARK_BLUE; r.font.name = "Arial"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Portfolio Risk Analytics & Signal Intelligence Report")
    r.font.size = Pt(13); r.font.color.rgb = BLUE; r.font.name = "Arial"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph_border_bottom(p)
    r = p.add_run(TODAY)
    r.font.size = Pt(11); r.font.color.rgb = GRAY; r.font.name = "Arial"

    spacer(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("20-Stock Equal-Weight Portfolio  |  S&P 500 Large-Cap Universe")
    r.font.size = Pt(10); r.font.color.rgb = GRAY; r.font.name = "Arial"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Analysis Period: January 2019 — December 2024")
    r.font.size = Pt(10); r.font.color.rgb = GRAY; r.font.name = "Arial"

    spacer(doc)

    # Cover KPI strip
    cover_kpis = [
        ("Ann. Return",   f"{port.get('ann_return_pct','—')}%"),
        ("Sharpe Ratio",  port.get("sharpe", "—")),
        ("Max Drawdown",  f"{port.get('max_drawdown_pct','—')}%"),
        ("VaR (95%)",     f"{port.get('hist_var_95_pct','—')}%"),
    ]
    tbl = doc.add_table(rows=2, cols=4)
    tbl.style = "Table Grid"
    for ci, (label, value) in enumerate(cover_kpis):
        c_top = tbl.rows[0].cells[ci]
        c_bot = tbl.rows[1].cells[ci]
        set_cell_bg(c_top, "1565C0"); set_cell_border(c_top, "0D3C6E")
        set_cell_bg(c_bot, "1565C0"); set_cell_border(c_bot, "0D3C6E")
        p_top = c_top.paragraphs[0]
        p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_top.add_run(label)
        r.font.size = Pt(8); r.font.name = "Arial"; r.font.color.rgb = RGBColor(0xCC,0xE5,0xFF)
        p_bot = c_bot.paragraphs[0]
        p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_bot.add_run(str(value))
        r.bold = True; r.font.size = Pt(16)
        r.font.name = "Arial"; r.font.color.rgb = WHITE

    doc.add_page_break()

    # ── SECTION 1: EXECUTIVE SUMMARY ──────────────────────────────────────────
    heading1(doc, "1. Executive Summary")
    body(doc, (
        "This report presents a comprehensive quantitative analysis of a 20-stock "
        "equal-weight portfolio drawn from five S&P 500 sectors: Technology, Financials, "
        "Healthcare, Energy, and Consumer Discretionary. The analysis covers January 2019 "
        "to December 2024, encompassing the COVID-19 market shock, the 2022 rate-hiking "
        "cycle, and the subsequent recovery."
    ))
    spacer(doc)
    body(doc, (
        "A machine learning signal model was developed to predict next-day return direction, "
        "validated on a strictly out-of-sample held-out test period using walk-forward "
        "methodology to prevent data leakage."
    ))
    spacer(doc)

    heading2(doc, "Key Findings")
    bullet(doc, f"Portfolio annualised return: {port.get('ann_return_pct','—')}%  |  Volatility: {port.get('ann_vol_pct','—')}%")
    bullet(doc, f"Sharpe Ratio: {port.get('sharpe','—')}  |  Sortino Ratio: {port.get('sortino','—')}")
    bullet(doc, f"Maximum Drawdown: {port.get('max_drawdown_pct','—')}%  |  Beta vs S&P 500: {port.get('beta_vs_sp500','—')}")
    bullet(doc, f"1-Day 95% VaR on $1M portfolio: {port.get('dollar_var_1M','—')}")
    bullet(doc, f"Best ML model: {best_ml.get('model','—')}  |  ROC-AUC: {best_ml.get('roc_auc','—')}  |  Accuracy: {best_ml.get('accuracy','—')}")
    if not top_feats.empty:
        bullet(doc, f"Top predictive feature: {top_feats.index[0]}")
    spacer(doc)

    # ── SECTION 2: PORTFOLIO RISK METRICS ────────────────────────────────────
    heading1(doc, "2. Portfolio Risk Metrics")
    body(doc, (
        "All metrics are annualised where applicable (252 trading days). "
        "Risk-free rate assumed at 5.0% per annum."
    ))
    spacer(doc)

    kpi_table(doc, [
        ("Annualised Return",           f"{port.get('ann_return_pct','—')}%"),
        ("Annualised Volatility",        f"{port.get('ann_vol_pct','—')}%"),
        ("Sharpe Ratio",                 port.get("sharpe","—")),
        ("Sortino Ratio",                port.get("sortino","—")),
        ("Maximum Drawdown",             f"{port.get('max_drawdown_pct','—')}%"),
        ("Historical VaR (95%, 1-Day)", f"{port.get('hist_var_95_pct','—')}%"),
        ("CVaR / Expected Shortfall",   f"{port.get('hist_cvar_95_pct','—')}%"),
        ("Dollar VaR ($1M Notional)",    port.get("dollar_var_1M","—")),
        ("Beta vs S&P 500",              port.get("beta_vs_sp500","—")),
    ])
    spacer(doc)

    heading2(doc, "Value at Risk Methodology")
    bullet(doc, "Historical Simulation VaR: Sorts actual observed returns, reads the loss at the 5th percentile. No distributional assumptions — most robust during regime changes.")
    bullet(doc, "Parametric VaR: Assumes normally distributed returns. Fast and widely used in internal bank models. Underestimates tail risk during stress periods.")
    bullet(doc, "Monte Carlo VaR: Simulates 10,000 return paths. Most flexible — can be extended to fat-tailed or regime-switching distributions.")
    spacer(doc)

    # ── SECTION 3: TOP STOCK PICKS ────────────────────────────────────────────
    heading1(doc, "3. Top Stock Picks by Risk-Adjusted Return")
    body(doc, "The following stocks ranked highest on Sharpe Ratio — strongest excess return per unit of total risk.")
    spacer(doc)

    if not top3.empty:
        cols  = ["ann_return_pct","ann_vol_pct","sharpe","sortino","max_drawdown_pct","beta"]
        cols  = [c for c in cols if c in top3.columns]
        hdrs  = ["Ticker","Ann Ret %","Ann Vol %","Sharpe","Sortino","Max DD %","Beta"][:len(cols)+1]
        rows  = [[idx] + [round(top3.loc[idx, c], 3) if c in top3.columns else "—" for c in cols]
                 for idx in top3.index]
        widths = [0.8] + [0.85] * len(cols)
        metrics_table(doc, hdrs, rows, widths)
    else:
        body(doc, "Run Module 2 (05_risk_analytics.py) to generate stock metrics.", italic=True, color=GRAY)
    spacer(doc)

    body(doc,
        "Note: Sortino Ratio penalises only downside volatility and is preferred "
        "for asymmetric return distributions.",
        italic=True, color=GRAY
    )
    spacer(doc)

    # ── SECTION 4: ML SIGNAL MODEL ────────────────────────────────────────────
    heading1(doc, "4. Machine Learning Signal Model")
    body(doc, (
        "A supervised classification model was trained to predict next-day return "
        "direction (Up/Down). Validated using strict walk-forward methodology — "
        "test set consists exclusively of dates after the training cutoff."
    ))
    spacer(doc)

    heading2(doc, "Model Architecture")
    bullet(doc, "Random Forest: 300 trees, max depth 6, balanced class weights, min 20 samples per leaf.")
    bullet(doc, "XGBoost: Gradient boosting, learning rate 0.05, column/row subsampling for regularisation.")
    bullet(doc, "Cross-Validation: 5-fold TimeSeriesSplit — each fold tests on a later window than its training set.")
    spacer(doc)

    heading2(doc, "Model Performance")
    if not ml_df.empty:
        ml_cols  = ["model","roc_auc","accuracy","precision_up","recall_up","f1_up","avg_precision"]
        ml_cols  = [c for c in ml_cols if c in ml_df.columns]
        ml_hdrs  = ["Model","ROC-AUC","Accuracy","Prec (Up)","Recall (Up)","F1 (Up)","Avg Prec"][:len(ml_cols)]
        ml_rows  = [[str(ml_df.loc[i, c]) for c in ml_cols] for i in ml_df.index]
        widths   = [1.3] + [0.78] * (len(ml_cols) - 1)
        metrics_table(doc, ml_hdrs, ml_rows, widths)
    else:
        body(doc, "Run Module 3 (06_ml_model.py) to generate ML evaluation results.", italic=True, color=GRAY)
    spacer(doc)

    body(doc,
        "Interpretation: ROC-AUC > 0.53 is meaningful for daily equity prediction. "
        "Markets are efficient by design — any edge is modest but probabilistically valuable "
        "when applied consistently across a large universe of stocks.",
    )
    spacer(doc)

    heading2(doc, "Top Predictive Features")
    body(doc, "Assessed using permutation importance — measures AUC decrease when each feature is shuffled.")
    spacer(doc)
    if not top_feats.empty:
        for i, (feat, row) in enumerate(top_feats.iterrows(), 1):
            imp = row.get("permutation_importance", "—")
            bullet(doc, f"{i}. {feat}  —  permutation importance: {round(float(imp), 4) if imp != '—' else '—'}")
    else:
        body(doc, "Run Module 3 to generate feature importance.", italic=True, color=GRAY)
    spacer(doc)

    # ── SECTION 5: RISK OBSERVATIONS ─────────────────────────────────────────
    heading1(doc, "5. Risk Observations & Recommendations")

    heading2(doc, "Concentration & Correlation Risk")
    body(doc, (
        "The correlation matrix reveals meaningful positive correlations within sectors, "
        "particularly Technology and Financials. During stress events, intra-sector "
        "correlations spike toward 1.0, reducing diversification benefit."
    ))
    bullet(doc, "Recommendation: Reduce Technology weight below 20% given its above-average beta (typically > 1.2).")
    bullet(doc, "Recommendation: Add genuinely uncorrelated assets (short-duration bonds, commodities) to improve portfolio efficiency.")
    spacer(doc)

    heading2(doc, "Drawdown & Tail Risk")
    bullet(doc, "The spread between parametric and historical VaR measures fat-tail exposure. A large spread indicates leptokurtic returns — more extreme losses than normality predicts.")
    bullet(doc, "CVaR (Expected Shortfall) is required under Basel III. Banks use CVaR alongside VaR for complete tail risk characterisation.")
    spacer(doc)

    heading2(doc, "Signal Model Deployment Considerations")
    bullet(doc, "The ML signal should be one input among several — combine with fundamental analysis and macro regime filters.")
    bullet(doc, "Monitor model performance monthly. Market regime changes can degrade performance and require retraining.")
    bullet(doc, "Transaction costs, market impact, and short-selling constraints are not modelled here.")
    spacer(doc)

    # ── SECTION 6: METHODOLOGY ────────────────────────────────────────────────
    heading1(doc, "6. Methodology & Data Sources")

    heading2(doc, "Data Sources")
    bullet(doc, "Price Data: Yahoo Finance via yfinance (adjusted close, split/dividend adjusted).")
    bullet(doc, "Macro Data: FRED — VIX, Fed Funds Rate, 10Y Treasury Yield, CPI, Unemployment, GDP Growth.")
    bullet(doc, "Benchmark: S&P 500 Index (^GSPC).")
    bullet(doc, "Analysis Period: January 2019 — December 2024 (~1,500 trading days).")
    spacer(doc)

    heading2(doc, "Technical Stack")
    bullet(doc, "Data Engineering: Python, Pandas, SQLAlchemy, SQLite")
    bullet(doc, "Risk Analytics: NumPy, SciPy (VaR, CVaR, Sharpe, Sortino, Beta, Drawdown)")
    bullet(doc, "Machine Learning: scikit-learn (Random Forest, TimeSeriesSplit CV), XGBoost")
    bullet(doc, "Visualisation: Matplotlib, Power BI / Tableau")
    bullet(doc, "Version Control: Git / GitHub")
    spacer(doc)

    body(doc,
        "This report was produced as a student data analysis project demonstrating "
        "end-to-end quantitative research capability. For educational purposes only — "
        "not investment advice.",
        italic=True, color=GRAY
    )

    # ── Save ──────────────────────────────────────────────────────────────────
    out_path = f"{OUT_DIR}/investment_memo.docx"
    doc.save(out_path)
    print(f"\n  Investment memo saved: {out_path}")
    print("  Open it in Microsoft Word or Google Docs.\n")


if __name__ == "__main__":
    print("\n" + "="*50)
    print("  MODULE 4b: INVESTMENT MEMO (Python)")
    print("="*50)
    build_memo()
